from __future__ import annotations
import sys as _public_sys
from pathlib import Path as _PublicPath
_public_root = next(p for p in _PublicPath(__file__).resolve().parents if (p / "common" / "public_paths.py").is_file())
if str(_public_root) not in _public_sys.path:
    _public_sys.path.insert(0, str(_public_root))
from common.public_paths import cohort_workbook, code_root, data_root, explainability_results, fusion_results, publication_results


import sys as _cohort_sys
from pathlib import Path as _CohortPath
_cohort_root = next(p for p in _CohortPath(__file__).resolve().parents if (p / "common" / "manuscript_cohorts.py").is_file())
if str(_cohort_root) not in _cohort_sys.path:
    _cohort_sys.path.insert(0, str(_cohort_root))
from common.manuscript_cohorts import split_center_a_indices, validate_cohort, validate_base_cohorts, split_metadata

import math
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import pandas as pd
from sklearn.model_selection import train_test_split
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HERE = Path(__file__).resolve()
V4_ROOT = HERE.parent.parent
FINAL_ROOT = V4_ROOT.parent
PROJECT_ROOT = data_root()
OUT = data_root() / "publication_results" / "Final_tables"
SOURCE_TABLES = OUT / "source_tables"

V1 = data_root() / "fusion_results" / "tables"
V2 = data_root() / "explainability_results" / "tables"
V4_TABLES = publication_results() / "tables"

BLACK = "000000"

def _research_root() -> Path:
    return data_root()


INPUT_WORKBOOK = cohort_workbook()
MASTER_SEED = 20260902

BASELINE_FEATURES = (
    ("WBC", "WBC (10^9/L)"),
    ("PLT", "PLT (10^9/L)"),
    ("LDH", "LDH (U/L)"),
    ("PCT", "PCT (%)"),
    ("EO#", "EO# (10^9/L)"),
    ("EO%", "EO% (%)"),
    ("ALB", "ALB (g/L)"),
    ("TP", "TP (g/L)"),
    ("GLB", "GLB (g/L)"),
    ("TC", "TC (mmol/L)"),
    ("LDL-C", "LDL-C (mmol/L)"),
    ("MCH", "MCH (pg)"),
    ("MCHC", "MCHC (g/L)"),
    ("SIRI", "SIRI"),
    ("LMR", "LMR"),
    ("GAR", "GAR"),
    ("PNI", "PNI"),
    ("HALP", "HALP"),
)

MODEL_NAMES = {
    "logistic_regression": "Logistic regression",
    "gam": "GAM",
    "knn": "KNN",
    "gaussian_nb": "Gaussian NB",
    "decision_tree": "Decision tree",
    "random_forest": "Random forest",
    "extra_trees": "Extra trees",
    "gbdt": "GBDT",
    "xgboost": "XGBoost",
    "lightgbm": "LightGBM",
    "adaboost": "AdaBoost",
    "rotation_forest": "Rotation forest",
    "mlp": "MLP",
    "rbf_svm": "RBF-SVM",
}


def fnum(value: object, digits: int = 3) -> str:
    if value is None or pd.isna(value):
        return "Not applicable"
    return f"{float(value):.{digits}f}"


def fci(value: object, low: object, high: object, digits: int = 3) -> str:
    if value is None or pd.isna(value):
        return "Not applicable"
    return f"{float(value):.{digits}f} ({float(low):.{digits}f}-{float(high):.{digits}f})"


def fdelta(value: object, digits: int = 3) -> str:
    if value is None or pd.isna(value):
        return "Not applicable"
    return f"{float(value):+.{digits}f}"


def fdelta_ci(value: object, low: object, high: object, digits: int = 3) -> str:
    if value is None or pd.isna(value):
        return "Not applicable"
    return f"{fdelta(value, digits)} ({float(low):.{digits}f}-{float(high):.{digits}f})"


def ensure_columns(frame: pd.DataFrame, columns: Iterable[str], source: Path) -> None:
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise ValueError(f"Missing required columns in {source.name}: {missing}")


def layer_label(value: str) -> str:
    return {"I": "I", "I_M": "I + M", "I_M_N": "I + M + N"}[value]


def dataset_label(value: str) -> str:
    labels = {
        "A_dev_cv_clean": "A development (10-fold CV)",
        "A_holdout_clean": "A evaluation set",
        "B_external": "B external",
        "C_external": "C external",
    }
    return labels[value]


def source_csv(name: str, frame: pd.DataFrame) -> Path:
    path = SOURCE_TABLES / name
    frame.to_csv(path, index=False, encoding="utf-8-sig")
    return path


def format_median_iqr(values: pd.Series) -> str:
    numeric = pd.to_numeric(values, errors="coerce").dropna().to_numpy(dtype=float)
    if not len(numeric):
        return "Not available"
    q1, median, q3 = np.quantile(numeric, [0.25, 0.50, 0.75])
    return f"{median:.2f} ({q1:.2f}-{q3:.2f})"


def chi_square_survival_df3(statistic: float) -> float:
    if not np.isfinite(statistic) or statistic < 0:
        return float("nan")
    z = math.sqrt(statistic / 2.0)
    return max(0.0, min(1.0, math.erfc(z) + (2.0 / math.sqrt(math.pi)) * z * math.exp(-z * z)))


def kruskal_wallis_pvalue(groups: Sequence[pd.Series]) -> float:
    clean = [pd.to_numeric(group, errors="coerce").dropna().to_numpy(dtype=float) for group in groups]
    clean = [group for group in clean if len(group)]
    if len(clean) != 4:
        return float("nan")
    combined = np.concatenate(clean)
    n_total = len(combined)
    if n_total < 5:
        return float("nan")
    ranks = pd.Series(combined).rank(method="average").to_numpy(dtype=float)
    starts = np.cumsum([0] + [len(group) for group in clean])
    rank_sum = [float(ranks[starts[index]:starts[index + 1]].sum()) for index in range(len(clean))]
    statistic = (12.0 / (n_total * (n_total + 1.0))) * sum(
        (value * value) / len(group) for value, group in zip(rank_sum, clean)
    ) - 3.0 * (n_total + 1.0)
    tie_sizes = pd.Series(combined).value_counts(dropna=False).to_numpy(dtype=float)
    correction = 1.0 - np.sum(tie_sizes**3 - tie_sizes) / (n_total**3 - n_total)
    if correction <= 0:
        return float("nan")
    return chi_square_survival_df3(float(statistic / correction))


def holm_adjust(p_values: Sequence[float]) -> np.ndarray:
    p = np.asarray(p_values, dtype=float)
    order = np.argsort(p, kind="mergesort")
    adjusted = np.empty_like(p)
    running = 0.0
    for rank, index in enumerate(order):
        running = max(running, min(1.0, (len(p) - rank) * p[index]))
        adjusted[index] = running
    return adjusted


def format_pvalue(value: float) -> str:
    if not np.isfinite(value):
        return "Not available"
    if value < 0.0001:
        return "<0.0001"
    return f"{value:.4f}"


def split_center_a(a_base: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    y = pd.to_numeric(a_base["y_SCLC"], errors="raise").to_numpy(dtype=int)
    indices = np.arange(len(a_base), dtype=int)
    train_idx, holdout_idx = split_center_a_indices(a_base)
    development = a_base.iloc[np.sort(train_idx)].reset_index(drop=True)
    evaluation = a_base.iloc[np.sort(holdout_idx)].reset_index(drop=True)
    if set(development["record_id"].astype(str)) & set(evaluation["record_id"].astype(str)):
        raise RuntimeError("A development and evaluation record ids overlap")
    return development, evaluation


def cohort_count_table(a_development: pd.DataFrame, a_evaluation: pd.DataFrame, b_external: pd.DataFrame, c_external: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for dataset, frame in (
        ("A_development", a_development),
        ("A_holdout_clean", a_evaluation),
        ("B_external", b_external),
        ("C_external", c_external),
    ):
        rows.append({
            "dataset": dataset,
            "n": int(len(frame)),
            "positive_n": int(pd.to_numeric(frame["y_SCLC"], errors="raise").sum()),
        })
    return pd.DataFrame(rows)


def load_baseline_cohorts() -> list[tuple[str, pd.DataFrame]]:
    if not INPUT_WORKBOOK.exists():
        raise FileNotFoundError(f"Baseline-table source is unavailable: {INPUT_WORKBOOK}")

    a_base = validate_cohort(pd.read_excel(INPUT_WORKBOOK, sheet_name="A_base"), 'A')
    b_external = validate_cohort(pd.read_excel(INPUT_WORKBOOK, sheet_name="B_base"), 'B')
    c_external = validate_cohort(pd.read_excel(INPUT_WORKBOOK, sheet_name="C_base"), 'C')
    validate_base_cohorts({"A": a_base, "B": b_external, "C": c_external})
    a_development, a_evaluation = split_center_a(a_base)
    required = ["record_id", "y_SCLC", *[feature for feature, _ in BASELINE_FEATURES]]
    for frame in (a_development, a_evaluation, b_external, c_external):
        ensure_columns(frame, required, INPUT_WORKBOOK)

    expected = {
        "A development": (split_metadata()["train_n"], split_metadata()["train_positive_n"]),
        "A evaluation set": (split_metadata()["test_n"], split_metadata()["test_positive_n"]),
        "B external": (329, 67),
        "C external": (313, 70),
    }
    cohorts = [
        ("A development", a_development),
        ("A evaluation set", a_evaluation),
        ("B external", b_external),
        ("C external", c_external),
    ]
    for label, frame in cohorts:
        n_expected, positive_expected = expected[label]
        positive_n = int(pd.to_numeric(frame["y_SCLC"], errors="raise").sum())
        if len(frame) != n_expected or positive_n != positive_expected:
            raise ValueError(
                f"Frozen cohort mismatch for {label}: observed n={len(frame)}, positives={positive_n}; "
                f"expected n={n_expected}, positives={positive_expected}."
            )
    return cohorts


def build_baseline_frame() -> tuple[pd.DataFrame, pd.DataFrame]:
    cohorts = load_baseline_cohorts()
    p_rows: list[dict[str, object]] = []
    headers = [f"{label}\n(n = {len(frame)})" for label, frame in cohorts]
    rows: list[dict[str, str]] = []
    event_row: dict[str, str] = {"Characteristic": "SCLC, n (%)", "P value": ""}
    for header, (_, frame) in zip(headers, cohorts):
        event_n = int(pd.to_numeric(frame["y_SCLC"], errors="raise").sum())
        event_row[header] = f"{event_n} ({100.0 * event_n / len(frame):.1f}%)"
    rows.append(event_row)

    for feature, label in BASELINE_FEATURES:
        p_value = kruskal_wallis_pvalue([frame[feature] for _, frame in cohorts])
        p_rows.append({"feature": feature, "kruskal_wallis_p": p_value})
        row = {"Characteristic": label, "P value": format_pvalue(p_value)}
        for header, (_, frame) in zip(headers, cohorts):
            row[header] = format_median_iqr(frame[feature])
        rows.append(row)
    pvalues = pd.DataFrame(p_rows)
    pvalues["holm_adjusted_p"] = holm_adjust(pvalues["kruskal_wallis_p"].to_numpy(dtype=float))
    return pd.DataFrame(rows, columns=["Characteristic", *headers, "P value"]), pvalues


def build_final_predictor_missingness_frame() -> pd.DataFrame:
    feature_order = [
        "WBC", "PLT", "LDH", "PCT", "EO#", "EO%", "ALB", "TP", "GLB", "TC", "LDL-C",
        "MCH", "MCHC", "SIRI", "LMR", "GAR", "PNI", "HALP",
    ]
    layers = {
        "WBC": "I", "PLT": "I", "LDH": "I", "PCT": "I", "SIRI": "I",
        "EO#": "M", "EO%": "M", "LMR": "M", "GAR": "M",
        "ALB": "N", "TP": "N", "GLB": "N", "TC": "N", "LDL-C": "N",
        "MCH": "N", "MCHC": "N", "PNI": "N", "HALP": "N",
    }
    centre_sheets = (("A centre", "A_base"), ("B centre", "B_base"), ("C centre", "C_base"))
    raw: dict[str, pd.DataFrame] = {}
    for centre, sheet in centre_sheets:
        frame = pd.read_excel(INPUT_WORKBOOK, sheet_name=sheet)
        ensure_columns(frame, feature_order, INPUT_WORKBOOK)
        raw[centre] = frame

    rows: list[dict[str, str]] = []
    for feature in feature_order:
        row = {"Predictor": feature, "Layer": layers[feature]}
        for centre, _ in centre_sheets:
            frame = raw[centre]
            missing_n = int(frame[feature].isna().sum())
            denominator = len(frame)
            row[f"{centre}, n/N (%)"] = (
                f"{missing_n}/{denominator} ({100.0 * missing_n / denominator:.2f}%)"
            )
        rows.append(row)
    return pd.DataFrame(
        rows,
        columns=["Predictor", "Layer", *[f"{centre}, n/N (%)" for centre, _ in centre_sheets]],
    )


def migrate_source_table_names() -> None:
    migrations = (
        ("table3_fusion_external_validation.csv", "table4_fusion_external_validation.csv"),
        ("table2_fusion_model_performance_A_evaluation.csv", "table3_fusion_model_performance_A_evaluation.csv"),
        ("table1_individual_model_performance_A_evaluation.csv", "table2_individual_model_performance_A_evaluation.csv"),
    )
    for old_name, new_name in migrations:
        old_path = SOURCE_TABLES / old_name
        new_path = SOURCE_TABLES / new_name
        if old_path.exists() and not new_path.exists():
            old_path.rename(new_path)


def set_run_font(run, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_border_edge(cell, edge: str, size: int, color: str = BLACK) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    tag = qn(f"w:{edge}")
    element = borders.find(tag)
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), color)


def clear_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def apply_three_line_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            clear_cell_borders(cell)
    for cell in table.rows[0].cells:
        set_border_edge(cell, "top", 12)
        set_border_edge(cell, "bottom", 8)
    for cell in table.rows[-1].cells:
        set_border_edge(cell, "bottom", 12)


def set_cell_margins(cell, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def write_cell(cell, value: object, *, header: bool, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(str(value))
    set_run_font(
        run,
        7.4 if not header else 7.6,
        bold=header,
    )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def make_document(title: str, introductory_sentence: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(9)

    document.core_properties.title = title
    document.core_properties.author = "Research team"
    return document


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_with_next(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, 9.5, bold=True)


def add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("Notes. ")
    set_run_font(run, 7.4, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, 7.4)


def fitted_widths(frame: pd.DataFrame, leading: Sequence[float], total: float = 27.0) -> list[float]:
    remaining = len(frame.columns) - len(leading)
    return [*leading, *([(total - sum(leading)) / remaining] * remaining)]


def add_table(
    document: Document,
    frame: pd.DataFrame,
    widths: Sequence[float],
    numeric_columns: Sequence[str] = (),
) -> None:
    if len(widths) != len(frame.columns):
        raise ValueError("Column-width count does not match table columns.")
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False

    header = table.rows[0]
    set_repeat_table_header(header)
    for index, name in enumerate(frame.columns):
        cell = header.cells[index]
        write_cell(
            cell,
            name,
            header=True,
            align=WD_ALIGN_PARAGRAPH.CENTER if name in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT,
        )
        set_cell_width(cell, widths[index])

    for row_index, (_, row) in enumerate(frame.iterrows()):
        cells = table.add_row().cells
        for column_index, column in enumerate(frame.columns):
            cell = cells[column_index]
            write_cell(
                cell,
                row[column],
                header=False,
                align=WD_ALIGN_PARAGRAPH.CENTER if column in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT,
            )
            set_cell_width(cell, widths[column_index])
    apply_three_line_borders(table)


def add_page_break(document: Document) -> None:
    document.add_page_break()


def build_table_frames() -> dict[str, pd.DataFrame]:
    baseline, baseline_pvalues = build_baseline_frame()

    individual_path = data_root() / "publication_results" / "tables" / "table2_model_performance_ci.csv"
    individual_raw = pd.read_csv(individual_path)
    ensure_columns(
        individual_raw,
        [
            "dataset", "model", "display_name", "score_kind", "auc", "auc_ci_low", "auc_ci_high",
            "auprc", "auprc_ci_low", "auprc_ci_high", "brier", "brier_ci_low", "brier_ci_high",
        ],
        individual_path,
    )
    individual_source = individual_raw.loc[
        (individual_raw["dataset"] == "A_holdout_clean")
        & ~individual_raw["model"].str.contains("voting|stacking", regex=True, na=False)
    ].copy()
    individual_source["Model"] = individual_source["display_name"].replace({"RBF-SVM": "RBF-SVM"})
    individual_source["Score output"] = np.where(
        individual_source["score_kind"].eq("decision_function"), "Decision score", "Probability"
    )
    individual_source["AUC (95% CI)"] = individual_source.apply(
        lambda row: fci(row.auc, row.auc_ci_low, row.auc_ci_high), axis=1
    )
    individual_source["AUPRC (95% CI)"] = individual_source.apply(
        lambda row: fci(row.auprc, row.auprc_ci_low, row.auprc_ci_high), axis=1
    )
    individual_source["Brier score (95% CI)"] = individual_source.apply(
        lambda row: fci(row.brier, row.brier_ci_low, row.brier_ci_high), axis=1
    )
    display_columns = ["Model", "Score output", "AUC (95% CI)", "AUPRC (95% CI)", "Brier score (95% CI)"]
    for label, point, low, high in (
        ("Accuracy (95% CI)", "accuracy", "accuracy_ci_low", "accuracy_ci_high"),
        ("Sensitivity (95% CI)", "sensitivity", "sensitivity_ci_low", "sensitivity_ci_high"),
        ("Specificity (95% CI)", "specificity", "specificity_ci_low", "specificity_ci_high"),
        ("F1 (95% CI)", "f1", "f1_ci_low", "f1_ci_high"),
        ("PPV (95% CI)", "ppv", "ppv_ci_low", "ppv_ci_high"),
        ("NPV (95% CI)", "npv", "npv_ci_low", "npv_ci_high"),
    ):
        if {point, low, high}.issubset(individual_source.columns):
            individual_source[label] = individual_source.apply(lambda row, point=point, low=low, high=high: fci(row[point], row[low], row[high]), axis=1)
            display_columns.append(label)
    individual = individual_source.sort_values(["auc", "Model"], ascending=[False, True])[display_columns].reset_index(drop=True)
    individual.attrs["threshold"] = float(individual_source["threshold_used"].dropna().iloc[0]) if "threshold_used" in individual_source else float("nan")

    fusion_path = V1 / "model_metrics_with_ci.csv"
    fusion_raw = pd.read_csv(fusion_path)

    ensure_columns(
        fusion_raw,
        [
            "dataset", "model", "auc", "auc_ci_low", "auc_ci_high", "auprc",
            "auprc_ci_low", "auprc_ci_high", "brier", "brier_ci_low", "brier_ci_high",
        ],
        fusion_path,
    )
    fusion_order = [
        "weighted_voting_cv_auc",
        "soft_voting_native_probability",
        "stacking_in_sample_A_train",
    ]
    fusion_names = {
        "weighted_voting_cv_auc": "Weighted voting (Primary)",
        "soft_voting_native_probability": "Soft voting",
        "stacking_in_sample_A_train": "Stacking",
    }

    def fusion_frame(dataset: str) -> pd.DataFrame:
        raw = fusion_raw.loc[
            (fusion_raw["dataset"] == dataset) & fusion_raw["model"].isin(fusion_order)
        ].copy()
        raw["model"] = pd.Categorical(raw["model"], categories=fusion_order, ordered=True)
        raw = raw.sort_values("model")
        raw["Model"] = raw["model"].astype(str).map(fusion_names)
        raw["AUC (95% CI)"] = raw.apply(lambda row: fci(row.auc, row.auc_ci_low, row.auc_ci_high), axis=1)
        raw["AUPRC (95% CI)"] = raw.apply(
            lambda row: fci(row.auprc, row.auprc_ci_low, row.auprc_ci_high), axis=1
        )
        raw["Brier score (95% CI)"] = raw.apply(
            lambda row: fci(row.brier, row.brier_ci_low, row.brier_ci_high), axis=1
        )
        columns = ["Model", "AUC (95% CI)", "AUPRC (95% CI)", "Brier score (95% CI)"]
        for label, point in (
            ("Calibration intercept (95% CI)", "calibration_intercept"),
            ("Calibration slope (95% CI)", "calibration_slope"),
            ("ECE (95% CI)", "ece_10bin"),
        ):
            if {point, f"{point}_ci_low", f"{point}_ci_high"}.issubset(raw.columns):
                raw[label] = raw.apply(lambda row, point=point: fci(row[point], row[f"{point}_ci_low"], row[f"{point}_ci_high"]), axis=1)
                columns.append(label)
        return raw[columns].reset_index(drop=True)

    fusion_eval = fusion_frame("A_holdout_clean")

    threshold_path = V1 / "all_probability_model_threshold_metrics.csv"
    threshold_raw = pd.read_csv(threshold_path)
    threshold_columns = (
        ("Accuracy (95% CI)", "accuracy"), ("Sensitivity (95% CI)", "sensitivity"),
        ("Specificity (95% CI)", "specificity"), ("F1 (95% CI)", "f1"),
        ("PPV (95% CI)", "ppv"), ("NPV (95% CI)", "npv"),
    )
    ensure_columns(
        threshold_raw,
        ["dataset", "model", "threshold", *[f"{m}{suffix}" for _, m in threshold_columns for suffix in ("", "_ci_low", "_ci_high")]],
        threshold_path,
    )
    threshold_datasets = ["A_dev_cv_clean", "A_holdout_clean", "B_external", "C_external"]
    fusion_threshold = threshold_raw.loc[
        threshold_raw["dataset"].isin(threshold_datasets) & threshold_raw["model"].isin(fusion_order)
    ].copy()
    if len(fusion_threshold) != len(threshold_datasets) * len(fusion_order):
        raise RuntimeError("Fixed-threshold metrics are incomplete for the three fusion models")
    decision_threshold = float(fusion_threshold["threshold"].iloc[0])
    fusion_threshold["order"] = (
        fusion_threshold["dataset"].map({name: i for i, name in enumerate(threshold_datasets)}) * 10
        + fusion_threshold["model"].map({name: i for i, name in enumerate(fusion_order)})
    )
    fusion_threshold = fusion_threshold.sort_values("order")
    fusion_threshold["Dataset"] = fusion_threshold["dataset"].map(
        lambda code: "A development (10-fold CV)" if code == "A_dev_cv_clean" else dataset_label(code)
    )
    fusion_threshold["Model"] = fusion_threshold["model"].map(fusion_names)
    for label, metric in threshold_columns:
        fusion_threshold[label] = fusion_threshold.apply(
            lambda row, metric=metric: fci(row[metric], row[f"{metric}_ci_low"], row[f"{metric}_ci_high"]), axis=1
        )
    fusion_threshold_display = fusion_threshold[["Dataset", "Model", *[label for label, _ in threshold_columns]]].reset_index(drop=True)
    fusion_threshold_display.attrs["threshold"] = decision_threshold
    external_parts = []
    for code in ("B_external", "C_external"):
        part = fusion_frame(code)
        part.insert(0, "Dataset", dataset_label(code))
        external_parts.append(part)
    fusion_external = pd.concat(external_parts, ignore_index=True)

    a_base = validate_cohort(pd.read_excel(INPUT_WORKBOOK, sheet_name="A_base"), 'A')
    b_external = validate_cohort(pd.read_excel(INPUT_WORKBOOK, sheet_name="B_base"), 'B')
    c_external = validate_cohort(pd.read_excel(INPUT_WORKBOOK, sheet_name="C_base"), 'C')
    validate_base_cohorts({"A": a_base, "B": b_external, "C": c_external})
    a_development, a_evaluation = split_center_a(a_base)
    cohort_raw = cohort_count_table(
        a_development,
        a_evaluation,
        b_external,
        c_external,
    )
    ensure_columns(cohort_raw, ["dataset", "n", "positive_n"], INPUT_WORKBOOK)
    cohort_codes = [
        "A_development",
        "A_holdout_clean",
        "B_external",
        "C_external",
    ]
    cohort = cohort_raw.loc[cohort_raw["dataset"].isin(cohort_codes), ["dataset", "n", "positive_n"]].copy()
    cohort["Dataset"] = cohort["dataset"].map(
        {
            "A_development": "A development (10-fold CV)",
            "A_holdout_clean": "A evaluation set",
            "B_external": "B external",
            "C_external": "C external",
        }
    )
    cohort["Study role"] = [
        "Development",
        "Evaluation",
        "External validation",
        "External validation",
    ]
    cohort["Participants"] = cohort["n"].astype(int).astype(str)
    cohort["SCLC, n (%)"] = cohort.apply(
        lambda row: f"{int(row.positive_n)} ({100 * row.positive_n / row.n:.1f}%)", axis=1
    )
    cohort["Non-SCLC, n (%)"] = cohort.apply(
        lambda row: f"{int(row.n - row.positive_n)} ({100 * (row.n - row.positive_n) / row.n:.1f}%)",
        axis=1,
    )
    cohort = cohort[["Dataset", "Study role", "Participants", "SCLC, n (%)", "Non-SCLC, n (%)"]]

    dictionary_path = code_root() / "config" / "feature_dictionary.csv"
    feature_raw = pd.read_csv(dictionary_path)
    ensure_columns(feature_raw, ["feature", "feature_type", "primary_layer"], dictionary_path)
    feature_order = [
        "WBC", "PLT", "LDH", "PCT", "EO#", "EO%", "ALB", "TP", "GLB", "TC", "LDL-C",
        "MCH", "MCHC", "SIRI", "LMR", "GAR", "PNI", "HALP",
    ]
    units = {
        "WBC": "10^9/L", "PLT": "10^9/L", "LDH": "U/L", "PCT": "%", "EO#": "10^9/L",
        "EO%": "%", "ALB": "g/L", "TP": "g/L", "GLB": "g/L", "TC": "mmol/L",
        "LDL-C": "mmol/L", "MCH": "pg", "MCHC": "g/L", "SIRI": "10^9/L",
        "LMR": "Index", "GAR": "Index", "PNI": "Index", "HALP": "Index",
    }
    definitions = {
        "WBC": "White blood cell count",
        "PLT": "Platelet count",
        "LDH": "Lactate dehydrogenase",
        "PCT": "Plateletcrit",
        "EO#": "Absolute eosinophil count",
        "EO%": "Eosinophil percentage",
        "ALB": "Albumin",
        "TP": "Total protein",
        "GLB": "Globulin",
        "TC": "Total cholesterol",
        "LDL-C": "Low-density lipoprotein cholesterol",
        "MCH": "Mean corpuscular hemoglobin",
        "MCHC": "Mean corpuscular hemoglobin concentration",
        "SIRI": "Systemic inflammation response index",
        "LMR": "Lymphocyte-to-monocyte ratio",
        "GAR": "Glucose-to-albumin ratio",
        "PNI": "Prognostic nutritional index",
        "HALP": "Hemoglobin, albumin, lymphocyte, and platelet index",
    }
    formulas = {
        "SIRI": "NEUT# x MONO# / LYMPH#",
        "LMR": "LYMPH# / MONO#",
        "GAR": "GLU / ALB",
        "PNI": "ALB + 5 x LYMPH#",
        "HALP": "HGB x ALB x LYMPH# / PLT",
    }
    feature_raw["order"] = feature_raw["feature"].map({name: i for i, name in enumerate(feature_order)})
    predictors = feature_raw.sort_values("order").copy()
    predictors["Predictor"] = predictors["feature"]
    predictors["Layer"] = predictors["primary_layer"]
    predictors["Type"] = predictors["feature_type"].map({"direct": "Direct", "composite": "Composite"})
    predictors["Unit"] = predictors["feature"].map(units)
    predictors["Definition or formula"] = predictors["feature"].map(
        lambda value: formulas[value] if value in formulas else definitions[value]
    )
    predictors = predictors[["Predictor", "Layer", "Type", "Unit", "Definition or formula"]]


    layer_path = V2 / "weighted_voting_layer_metrics_ci_harmonised.csv"
    layer_raw = pd.read_csv(layer_path)
    ensure_columns(
        layer_raw,
        [
            "layer", "dataset", "auc", "auc_ci_low", "auc_ci_high", "auprc",
            "auprc_ci_low", "auprc_ci_high", "brier", "brier_ci_low", "brier_ci_high",
        ],
        layer_path,
    )
    layer_dataset_codes = [
        "A_holdout_clean",
        "B_external",
        "C_external",
    ]
    layer_order = ["I", "I_M", "I_M_N"]
    layer_display = layer_raw.loc[
        layer_raw["dataset"].isin(layer_dataset_codes) & layer_raw["layer"].isin(layer_order)
    ].copy()
    layer_display["order"] = layer_display["dataset"].map({name: i for i, name in enumerate(layer_dataset_codes)}) * 10
    layer_display["order"] += layer_display["layer"].map({name: i for i, name in enumerate(layer_order)})
    layer_display = layer_display.sort_values("order")
    layer_display["Dataset"] = layer_display["dataset"].map(dataset_label)
    layer_display["Layer"] = layer_display["layer"].map(layer_label)
    layer_display["AUC (95% CI)"] = layer_display.apply(
        lambda row: fci(row.auc, row.auc_ci_low, row.auc_ci_high), axis=1
    )
    layer_display["AUPRC (95% CI)"] = layer_display.apply(
        lambda row: fci(row.auprc, row.auprc_ci_low, row.auprc_ci_high), axis=1
    )
    layer_display["Brier score (95% CI)"] = layer_display.apply(
        lambda row: fci(row.brier, row.brier_ci_low, row.brier_ci_high), axis=1
    )
    layer_display = layer_display[
        ["Dataset", "Layer", "AUC (95% CI)", "AUPRC (95% CI)", "Brier score (95% CI)"]
    ].reset_index(drop=True)

    increment_path = V1 / "weighted_voting_layer_increment_ci.csv"
    delong_path = V1 / "weighted_voting_layer_delong_holm.csv"
    increments = pd.read_csv(increment_path)
    delong = pd.read_csv(delong_path)
    ensure_columns(increments, ["comparison", "metric", "increment", "ci_low", "ci_high"], increment_path)
    ensure_columns(delong, ["comparison", "p_raw", "p_holm"], delong_path)
    inc_a = increments.loc[increments["dataset"] == "A_dev_cv_clean"].copy()
    delong_a = delong.loc[delong["dataset"] == "A_dev_cv_clean"].copy()
    comparison_order = [
        ("I", "I_M", "I to I + M"),
        ("I_M", "I_M_N", "I + M to I + M + N"),
        ("I", "I_M_N", "I to I + M + N"),
    ]
    increment_rows = []
    for lower_layer, upper_layer, comparison_label in comparison_order:
        selected = inc_a.loc[
            (inc_a["lower_layer"] == lower_layer) & (inc_a["upper_layer"] == upper_layer)
        ].set_index("metric")
        test = delong_a.loc[
            (delong_a["lower_layer"] == lower_layer) & (delong_a["upper_layer"] == upper_layer)
        ].iloc[0]
        increment_rows.append(
            {
                "Comparison": comparison_label,
                "AUC change (95% CI)": fdelta_ci(
                    selected.loc["auc", "increment"], selected.loc["auc", "ci_low"], selected.loc["auc", "ci_high"]
                ),
                "AUPRC change (95% CI)": fdelta_ci(
                    selected.loc["auprc", "increment"], selected.loc["auprc", "ci_low"], selected.loc["auprc", "ci_high"]
                ),
                "Brier change (95% CI)": fdelta_ci(
                    selected.loc["brier", "increment"], selected.loc["brier", "ci_low"], selected.loc["brier", "ci_high"]
                ),
                "DeLong P": f"{float(test.p_raw):.4f}",
                "Holm-adjusted P": f"{float(test.p_holm):.4f}",
            }
        )
    increments_display = pd.DataFrame(increment_rows)

    selection_path = V4_TABLES / "TableS8_development_cv_selection_evidence.csv"
    selection = pd.read_csv(selection_path)
    ensure_columns(
        selection,
        [
            "model", "auc", "auc_ci_low", "auc_ci_high", "auprc", "auprc_ci_low",
            "auprc_ci_high", "brier", "brier_ci_low", "brier_ci_high",
        ],
        selection_path,
    )
    selection["Model"] = selection["model"].replace({"Weighted voting": "Weighted voting (Primary)"})
    selection["AUC (95% CI)"] = selection.apply(
        lambda row: fci(row.auc, row.auc_ci_low, row.auc_ci_high), axis=1
    )
    selection["AUPRC (95% CI)"] = selection.apply(
        lambda row: fci(row.auprc, row.auprc_ci_low, row.auprc_ci_high), axis=1
    )
    selection["Brier score (95% CI)"] = selection.apply(
        lambda row: fci(row.brier, row.brier_ci_low, row.brier_ci_high), axis=1
    )
    selection_display = selection[
        ["Model", "AUC (95% CI)", "AUPRC (95% CI)", "Brier score (95% CI)"]
    ].copy()

    fusion_construction = pd.DataFrame(
        [
            {
                "Fusion model": "Weighted voting (Primary)",
                "Base learners": "13",
                "Aggregation": "AUC-normalized weighted mean",
                "Fixed input rule": "Development CV AUC",
            },
            {
                "Fusion model": "Soft voting",
                "Base learners": "13",
                "Aggregation": "Equal-weight mean",
                "Fixed input rule": "Native probabilities",
            },
            {
                "Fusion model": "Stacking",
                "Base learners": "13",
                "Aggregation": "Logistic-regression meta-learner",
                "Fixed input rule": "Native probabilities",
            },
        ]
    )

    weights_path = V1 / "weighted_voting_fixed_weights.csv"
    weights = pd.read_csv(weights_path)
    ensure_columns(weights, ["model", "source_cv_auc", "normalized_weight"], weights_path)
    if "layer" in weights.columns:
        weights = weights.loc[weights["layer"] == "I_M_N"].reset_index(drop=True)
    weights["Base model"] = weights["model"].map(MODEL_NAMES)
    weights["Development CV AUC"] = weights["source_cv_auc"].map(fnum)
    weights["Normalized weight"] = weights["normalized_weight"].map(lambda value: f"{100 * value:.2f}%")
    weights_display = weights[["Base model", "Development CV AUC", "Normalized weight"]].copy()

    shap_path = V1 / "F11_shap_global_importance.csv"
    shap = pd.read_csv(shap_path)
    ensure_columns(shap, ["rank", "feature", "feature_group", "mean_abs_shap", "mean_shap", "shap_sd"], shap_path)
    shap_display = pd.DataFrame(
        {
            "Rank": shap["rank"].astype(int),
            "Feature": shap["feature"],
            "Layer": shap["feature_group"],
            "Mean absolute attribution": shap["mean_abs_shap"].map(lambda value: fnum(value, 4)),
            "Mean signed attribution": shap["mean_shap"].map(lambda value: fnum(value, 4)),
            "Attribution SD": shap["shap_sd"].map(lambda value: fnum(value, 4)),
        }
    )

    interaction_path = V2 / "weighted_sii_global_pair_ranking.csv"
    interactions = pd.read_csv(interaction_path)
    ensure_columns(
        interactions,
        [
            "rank", "pair", "mean_abs_sii", "median_abs_sii", "mean_signed_sii",
            "sd_signed_sii_across_cases", "mean_seed_sd",
        ],
        interaction_path,
    )
    interaction_display = pd.DataFrame(
        {
            "Rank": interactions.loc[interactions["rank"] <= 10, "rank"].astype(int),
            "Feature pair": interactions.loc[interactions["rank"] <= 10, "pair"],
            "Mean absolute SII": interactions.loc[interactions["rank"] <= 10, "mean_abs_sii"].map(
                lambda value: fnum(value, 4)
            ),
            "Median absolute SII": interactions.loc[interactions["rank"] <= 10, "median_abs_sii"].map(
                lambda value: fnum(value, 4)
            ),
            "Mean signed SII": interactions.loc[interactions["rank"] <= 10, "mean_signed_sii"].map(
                lambda value: fnum(value, 4)
            ),
            "SD of signed SII": interactions.loc[
                interactions["rank"] <= 10, "sd_signed_sii_across_cases"
            ].map(lambda value: fnum(value, 4)),
            "Mean seed SD": interactions.loc[interactions["rank"] <= 10, "mean_seed_sd"].map(
                lambda value: fnum(value, 4)
            ),
        }
    )
    missingness_display = build_final_predictor_missingness_frame()

    return {
        "table1_baseline_laboratory_characteristics": baseline,
        "table2_individual_model_performance_A_evaluation": individual,
        "table3_fusion_model_performance_A_evaluation": fusion_eval,
        "table4_fusion_external_validation": fusion_external,
        "table5_fusion_fixed_threshold_classification": fusion_threshold_display,
        "table1_kruskal_wallis_holm": baseline_pvalues,
        "tableS1_predictor_definitions": predictors,
        "tableS2_analysis_cohort_composition": cohort,
        "tableS3_layerwise_weighted_voting_performance": layer_display,
        "tableS4_layer_increment_delong_holm": increments_display,
        "tableS5a_fusion_selection_evidence": selection_display,
        "tableS5b_fusion_model_construction": fusion_construction,
        "tableS5c_weighted_voting_base_model_weights": weights_display,
        "tableS6_global_attribution_importance": shap_display,
        "tableS7_pairwise_interaction_effects": interaction_display,
        "tableS8_final_predictor_missingness": missingness_display,
    }


def build_main_document(frames: dict[str, pd.DataFrame]) -> Path:
    document = make_document(
        "Main Tables",
        "Tables are presented in the same order as the final figure narrative.",
    )
    add_caption(document, "Table 1. Baseline laboratory characteristics of the analytic cohorts")
    add_table(
        document,
        frames["table1_baseline_laboratory_characteristics"],
        [4.1, 4.8, 4.8, 4.8, 4.8, 1.9],
        numeric_columns=[
            f"A development\n(n = {split_metadata()['train_n']})", f"A evaluation set\n(n = {split_metadata()['test_n']})",
            "B external\n(n = 329)", "C external\n(n = 313)", "P value",
        ],
    )
    add_note(
        document,
        "Continuous variables are reported as median (Q1-Q3). P values are from four-group "
        "Kruskal-Wallis tests without multiplicity adjustment; after Holm adjustment across the "
        f"{len(frames['table1_kruskal_wallis_holm'])} comparisons the smallest adjusted P value was "
        f"{frames['table1_kruskal_wallis_holm']['holm_adjusted_p'].min():.3f}. "
        "SCLC, small-cell lung cancer; IQR, interquartile range.",
    )
    add_page_break(document)

    add_caption(document, "Table 2. Individual-model performance on the A evaluation set")
    add_table(
        document,
        frames["table2_individual_model_performance_A_evaluation"],
        fitted_widths(frames["table2_individual_model_performance_A_evaluation"], [3.2, 2.1]),
        numeric_columns=list(frames["table2_individual_model_performance_A_evaluation"].columns[2:]),
    )
    add_note(
        document,
        "Intervals are outcome-stratified bootstrap 95% confidence intervals. "
        f"Threshold-dependent metrics use the development CV Youden threshold {frames['table2_individual_model_performance_A_evaluation'].attrs['threshold']:.4f}. "
        "RBF-SVM generated a decision score; probability-based Brier and threshold metrics are not applicable.",
    )
    add_page_break(document)

    add_caption(document, "Table 3. Fusion-model performance on the A evaluation set")
    add_table(
        document,
        frames["table3_fusion_model_performance_A_evaluation"],
        fitted_widths(frames["table3_fusion_model_performance_A_evaluation"], [4.5]),
        numeric_columns=list(frames["table3_fusion_model_performance_A_evaluation"].columns[1:]),
    )
    add_note(
        document,
        "Weighted voting was pre-specified as the primary model according to A development 10-fold CV selection evidence.",
    )
    add_page_break(document)

    add_caption(document, "Table 4. External validation performance of fusion models")
    add_table(
        document,
        frames["table4_fusion_external_validation"],
        fitted_widths(frames["table4_fusion_external_validation"], [3.0, 4.2]),
        numeric_columns=list(frames["table4_fusion_external_validation"].columns[2:]),
    )
    add_note(
        document,
        "The fixed model order is retained across B and C external validation. "
        "All intervals are outcome-stratified bootstrap 95% confidence intervals.",
    )
    add_page_break(document)

    add_caption(document, "Table 5. Fixed-threshold classification performance of the fusion models")
    add_table(
        document,
        frames["table5_fusion_fixed_threshold_classification"],
        fitted_widths(frames["table5_fusion_fixed_threshold_classification"], [3.4, 3.4]),
        numeric_columns=[
            "Accuracy (95% CI)", "Sensitivity (95% CI)", "Specificity (95% CI)",
            "F1 (95% CI)", "PPV (95% CI)", "NPV (95% CI)",
        ],
    )
    add_note(
        document,
        f"The decision threshold {frames['table5_fusion_fixed_threshold_classification'].attrs['threshold']:.4f} is the maximum "
        "Youden index of the primary model's pooled development 10-fold CV predictions and is applied unchanged to "
        "every dataset and model; predictions at or above it are classified as SCLC.",
    )
    path = OUT / "Main_Tables.docx"
    document.save(path)
    return path


def build_supplementary_document(frames: dict[str, pd.DataFrame]) -> Path:
    document = make_document(
        "Supplementary Tables",
        "Supporting definitions, selection evidence, and model interpretation tables.",
    )

    add_caption(document, "Table S5. Final predictor set and composite-index definitions")
    add_table(
        document,
        frames["tableS1_predictor_definitions"],
        [2.0, 1.3, 1.5, 2.0, 11.3],
        numeric_columns=[],
    )
    add_note(
        document,
        "I, inflammation; M, immune; N, nutrition or metabolism. "
        "Hash marks identify absolute cell counts.",
    )
    add_page_break(document)

    add_caption(document, "Table S4. Analysis cohort composition")
    add_table(
        document,
        frames["tableS2_analysis_cohort_composition"],
        [5.5, 4.2, 3.0, 4.2, 4.2],
        numeric_columns=["Participants", "SCLC, n (%)", "Non-SCLC, n (%)"],
    )
    add_note(document, "SCLC, small-cell lung cancer.")
    add_page_break(document)

    add_caption(document, "Table S9. Weighted-voting performance across cumulative predictor layers")
    add_table(
        document,
        frames["tableS3_layerwise_weighted_voting_performance"],
        [3.8, 2.4, 5.2, 5.2, 5.2],
        numeric_columns=["AUC (95% CI)", "AUPRC (95% CI)", "Brier score (95% CI)"],
    )
    add_note(
        document,
        "I, inflammation; M, immune; N, nutrition or metabolism. "
        "All intervals are outcome-stratified bootstrap 95% confidence intervals.",
    )
    add_page_break(document)

    add_caption(document, "Table S10. Incremental layer comparison in A development 10-fold CV")
    add_table(
        document,
        frames["tableS4_layer_increment_delong_holm"],
        [3.4, 4.2, 4.2, 4.2, 2.5, 2.7],
        numeric_columns=[
            "AUC change (95% CI)", "AUPRC change (95% CI)", "Brier change (95% CI)",
            "DeLong P", "Holm-adjusted P",
        ],
    )
    add_note(
        document,
        "Changes are calculated as upper minus lower layer. Confidence intervals are paired "
        "outcome-stratified bootstrap intervals. The three pre-specified AUC comparisons were "
        "evaluated with paired DeLong tests and Holm adjustment.",
    )
    add_page_break(document)

    add_caption(document, "Table S8. Fusion-model selection evidence in A development 10-fold CV")
    add_table(
        document,
        frames["tableS5a_fusion_selection_evidence"],
        [5.0, 5.3, 5.3, 5.3],
        numeric_columns=["AUC (95% CI)", "AUPRC (95% CI)", "Brier score (95% CI)"],
    )
    add_note(
        document,
        "Selection evidence is based on pooled out-of-fold patient-level predictions. "
        "Weighted voting was selected before evaluation and external validation.",
    )
    add_page_break(document)

    add_caption(document, "Table S3. Construction of the three fusion models")
    add_table(
        document,
        frames["tableS5b_fusion_model_construction"],
        [5.2, 2.7, 6.5, 6.5],
        numeric_columns=["Base learners"],
    )
    add_note(
        document,
        "All three fusion models used the same 13 base-model probability outputs; only the aggregation rule differed.",
    )
    add_page_break(document)

    add_caption(document, "Table S7. Fixed base-model weights for weighted voting")
    add_table(
        document,
        frames["tableS5c_weighted_voting_base_model_weights"],
        [8.5, 5.5, 5.5],
        numeric_columns=["Development CV AUC", "Normalized weight"],
    )
    add_note(document, "Weights are normalized from prespecified development CV AUC values.")
    add_page_break(document)

    add_caption(document, "Table S11. Global attribution ranking for the weighted-voting model")
    add_table(
        document,
        frames["tableS6_global_attribution_importance"],
        [1.5, 2.7, 1.4, 5.2, 5.2, 4.3],
        numeric_columns=[
            "Rank", "Mean absolute attribution", "Mean signed attribution", "Attribution SD",
        ],
    )
    add_note(
        document,
        "Attribution values summarize model-agnostic weighted-voting explanations; higher mean "
        "absolute attribution indicates greater global contribution.",
    )
    add_page_break(document)

    add_caption(document, "Table S12. Top estimated pairwise interaction effects for weighted voting")
    add_table(
        document,
        frames["tableS7_pairwise_interaction_effects"],
        [1.4, 4.0, 3.8, 3.8, 3.6, 3.7, 3.0],
        numeric_columns=[
            "Rank", "Mean absolute SII", "Median absolute SII", "Mean signed SII",
            "SD of signed SII", "Mean seed SD",
        ],
    )
    add_note(
        document,
        "SII, estimated pairwise interaction index. Feature pairs are ranked by mean absolute SII.",
    )
    add_page_break(document)

    add_caption(document, "Table S6. Pre-imputation missingness of the final predictor set in the raw centre cohorts")
    add_table(
        document,
        frames["tableS8_final_predictor_missingness"],
        [2.6, 1.4, 5.8, 5.8, 5.8],
        numeric_columns=["A centre, n/N (%)", "B centre, n/N (%)", "C centre, n/N (%)"],
    )
    add_note(
        document,
        "Denominators are raw source records before imputation and before downstream data-set construction "
        f"(A, n = {split_metadata()['source_n']}; B, n = 329; C, n = 313). Direct predictors therefore use the same source "
        "cohorts and denominators as Figure 1c. Composite indices use their stored values before imputation.",
    )
    path = OUT / "Supplementary_Tables.docx"
    document.save(path)
    return path


def build_index() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["Table 1", "Baseline laboratory characteristics of the analytic cohorts", "F01", "Main_Tables.docx"],
            ["Table 2", "Individual-model performance on the A evaluation set", "F02-F05", "Main_Tables.docx"],
            ["Table 3", "Fusion-model performance on the A evaluation set", "F06", "Main_Tables.docx"],
            ["Table 4", "External validation performance of fusion models", "F07", "Main_Tables.docx"],
            ["Table 5", "Fixed-threshold classification performance of the fusion models", "F06-F07", "Main_Tables.docx"],
            ["Table S5", "Final predictor set and composite-index definitions", "Fig1", "Supplementary_Tables.docx"],
            ["Table S4", "Analysis cohort composition", "Fig1", "Supplementary_Tables.docx"],
            ["Table S9", "Weighted-voting performance across cumulative predictor layers", "Fig6", "Supplementary_Tables.docx"],
            ["Table S10", "Incremental layer comparison in A development 10-fold CV", "Fig6", "Supplementary_Tables.docx"],
            ["Table S8", "Fusion-model selection evidence in A development 10-fold CV", "Fig4", "Supplementary_Tables.docx"],
            ["Table S3", "Construction of the three fusion models", "Fig4-Fig5", "Supplementary_Tables.docx"],
            ["Table S7", "Fixed base-model weights for weighted voting", "Fig4-Fig5", "Supplementary_Tables.docx"],
            ["Table S11", "Global attribution ranking for the weighted-voting model", "Fig7", "Supplementary_Tables.docx"],
            ["Table S12", "Top estimated pairwise interaction effects for weighted voting", "FigS8", "Supplementary_Tables.docx"],
            ["Table S6", "Pre-imputation missingness of the final predictor set in the raw centre cohorts", "Fig2", "Supplementary_Tables.docx"],
        ],
        columns=["Table", "Title", "Associated figures", "Document"],
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    SOURCE_TABLES.mkdir(parents=True, exist_ok=True)
    migrate_source_table_names()
    frames = build_table_frames()
    for name, frame in frames.items():
        source_csv(f"{name}.csv", frame)
    build_index().to_csv(OUT / "table_index.csv", index=False, encoding="utf-8-sig")
    main_path = build_main_document(frames)
    supplementary_path = build_supplementary_document(frames)
    print(f"Created {main_path}")
    print(f"Created {supplementary_path}")
    print(f"Created {len(frames)} source tables in {SOURCE_TABLES}")


if __name__ == "__main__":
    main()
